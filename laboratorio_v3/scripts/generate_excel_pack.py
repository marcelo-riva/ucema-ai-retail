from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PORTFOLIO_PACK_FILENAME = "NEXUS_RETAIL_LAB01_EJ01_PORTFOLIO_PACK_FINAL.xlsx"
SOURCE_PATH = ROOT / "public" / "templates" / PORTFOLIO_PACK_FILENAME
APP_PUBLIC_PATH = ROOT / "apps" / "web" / "public" / "templates" / PORTFOLIO_PACK_FILENAME

EXPECTED_SHEETS = [
    "99_VALIDACIONES",
    "00_CASO_NEGOCIO",
    "01_BASE_SKUS",
    "02_DICCIONARIO_DATOS",
    "03_GUIA_EXPLORACION",
    "04_HIPOTESIS_EQUIPO",
    "05_DECISIONES_SKU",
    "09_PROYECCION_90_DIAS",
    "06_SCOREBOARD_ALUMNO",
    "07_PLAN_90_DIAS",
    "08_OUTPUT_FINAL",
]

DECISION_HEADERS = [
    "sku_id",
    "sku_name",
    "category",
    "revenue_12m",
    "gross_margin_12m",
    "gross_margin_pct",
    "current_inventory_value",
    "ddi_current",
    "market_coverage",
    "decision_portfolio",
    "action_90_days",
    "decision_reason",
]


def validate_pack(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"No existe el pack final: {path}")

    workbook = load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise ValueError(f"Hojas inesperadas: {workbook.sheetnames}")

    decisions = workbook["05_DECISIONES_SKU"]
    decision_headers = [cell.value for cell in decisions[1][: len(DECISION_HEADERS)]]
    if decision_headers != DECISION_HEADERS:
        raise ValueError(f"Columnas iniciales inesperadas en 05_DECISIONES_SKU: {decision_headers}")

    projection = workbook["09_PROYECCION_90_DIAS"]
    projection_headers = [cell.value for cell in projection[5] if cell.value]
    if not any(str(header).startswith("baseline_") for header in projection_headers):
        raise ValueError("09_PROYECCION_90_DIAS no contiene columnas baseline_*")
    if not any(str(header).startswith("projected_") for header in projection_headers):
        raise ValueError("09_PROYECCION_90_DIAS no contiene columnas projected_*")


def main() -> None:
    validate_pack(SOURCE_PATH)
    APP_PUBLIC_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(SOURCE_PATH, APP_PUBLIC_PATH)
    print(f"Pack final validado: {SOURCE_PATH}")
    print(f"Copia para Next.js actualizada: {APP_PUBLIC_PATH}")


if __name__ == "__main__":
    main()
