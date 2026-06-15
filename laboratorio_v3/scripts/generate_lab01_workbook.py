from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = ROOT.parent
BASE_SKUS = REPO_ROOT / "data" / "Base SKUs.xlsx"
OUTPUT = ROOT / "public" / "templates" / "NEXUS_RETAIL_LAB01_WORKBOOK_COMPLETO.xlsx"
APP_OUTPUT = ROOT / "apps" / "web" / "public" / "templates" / "NEXUS_RETAIL_LAB01_WORKBOOK_COMPLETO.xlsx"

EXPECTED_SHEETS = [
    "00_CASO_NEGOCIO",
    "01_BASE_SKUS",
    "02_DICCIONARIO_DATOS",
    "03_GUIA_EXPLORACION",
    "04_DIAGNOSTICO_INICIAL",
    "05_DECISIONES_PORTFOLIO",
    "06_PRICING_DECISIONS",
    "07_FORECAST_90_DIAS",
    "08_INVENTORY_DECISIONS",
    "09_SCOREBOARD_ALUMNO",
    "10_PLAN_90_DIAS",
    "11_OUTPUT_FINAL",
    "12_CONTROL_STATUS",
]

HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FEF3C7")


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def write_rows(workbook: Workbook, title: str, rows: list[list[object]]) -> None:
    ws = workbook.create_sheet(title)
    for row in rows:
        ws.append(row)
    if rows:
        style_header(ws)
    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 14), 48)


def copy_base_skus(workbook: Workbook) -> None:
    if not BASE_SKUS.exists():
        raise FileNotFoundError(f"No existe la base SKU: {BASE_SKUS}")

    source_workbook = load_workbook(BASE_SKUS, data_only=False, read_only=True)
    source = source_workbook["Base"]
    ws = workbook.create_sheet("01_BASE_SKUS")

    for row in source.iter_rows(values_only=True):
        ws.append(list(row))

    ws.freeze_panes = "A8"
    for row in ws.iter_rows(min_row=1, max_row=7):
        for cell in row:
            cell.fill = NOTE_FILL
            cell.font = Font(bold=True)
    ws["A1"] = "NO EDITAR: esta hoja conserva la base historica M01-M12 del Laboratorio 1."
    ws.column_dimensions["A"].width = 18


def add_list_validation(ws, column: str, values: list[str], start_row: int = 2, end_row: int = 5000) -> None:
    validation = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"{column}{start_row}:{column}{end_row}")


def build_workbook() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)

    write_rows(
        workbook,
        "00_CASO_NEGOCIO",
        [
            ["campo", "contenido"],
            ["nombre_caso", "Nexus Retail: AI Revenue & Inventory Copilot"],
            ["contexto", "Nexus Retail tiene un portfolio amplio, ventas historicas M01-M12, costos, precios, stock, DDI y cobertura."],
            ["objetivo_laboratorio", "Explorar la base, decidir portfolio, definir pricing, proyectar 90 dias, optimizar inventario y cerrar un plan de captura de valor."],
            ["horizonte", "90 dias"],
            ["rol_alumno", "Equipo comercial y estrategico responsable de recomendar decisiones defendibles."],
            ["rol_ai_personal", "Asistente de analisis. La AI no decide por el equipo."],
            ["principio_operativo", "Si la plataforma falla, continuar desde este workbook con una AI personal."],
        ],
    )
    copy_base_skus(workbook)
    write_rows(
        workbook,
        "02_DICCIONARIO_DATOS",
        [
            ["campo", "descripcion", "por_que_importa"],
            ["Venta Neta", "Ventas valorizadas sin IVA.", "Permite priorizar productos por contribucion comercial."],
            ["PVP 1-12", "Precio historico mensual.", "Ayuda a revisar pricing y sensibilidad precio-volumen."],
            ["CU 1-12", "Costo unitario historico.", "Permite estimar margen y presion de costos."],
            ["Vol 1-12", "Unidades vendidas historicas.", "Permite mirar tendencia, estacionalidad y caidas."],
            ["DDI", "Dias de inventario disponible.", "DDI alto inmoviliza capital; DDI bajo puede generar quiebres."],
            ["Mercado %", "Cobertura o presencia comercial.", "Ayuda a identificar productos con rol estrategico."],
        ],
    )
    write_rows(
        workbook,
        "03_GUIA_EXPLORACION",
        [
            ["pregunta_de_analisis", "variables_sugeridas", "que_buscar", "hallazgo_del_equipo", "implicancia_para_la_decision"],
            ["Que categorias concentran ventas?", "Venta Neta, Familia, Departamento", "Concentracion de revenue", "", ""],
            ["Que categorias concentran margen?", "Venta, CU, PVP, margen estimado", "Categorias rentables o presionadas", "", ""],
            ["Donde hay mas stock o DDI?", "Inventario TOTAL, DDI", "Capital inmovilizado", "", ""],
            ["Que productos muestran caida reciente?", "Vol 10-12 vs Vol 1-3", "Tendencia negativa", "", ""],
            ["Que productos tienen margen negativo o bajo?", "PVP, CU, Venta", "Leakage o costo no trasladado", "", ""],
            ["Que preguntas debe investigar el equipo?", "Todas", "Dudas antes de decidir", "", ""],
        ],
    )
    write_rows(
        workbook,
        "04_DIAGNOSTICO_INICIAL",
        [
            ["campo", "respuesta_del_equipo"],
            ["diagnostico_inicial", ""],
            ["cinco_hallazgos_principales", ""],
            ["tres_preguntas_de_negocio", ""],
            ["supuestos_o_dudas", ""],
        ],
    )
    write_rows(
        workbook,
        "05_DECISIONES_PORTFOLIO",
        [
            ["sku_id", "sku_name", "category", "historical_revenue_12m", "historical_margin_12m", "current_inventory_value", "ddi", "market_coverage", "decision_portfolio", "action_90_days", "decision_reason", "priority", "commercial_risk", "ai_comment", "team_comment"],
            ["Completar en Ejercicio 1", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
    )
    portfolio_ws = workbook["05_DECISIONES_PORTFOLIO"]
    add_list_validation(portfolio_ws, "I", ["CORE", "REVIEW", "ELIMINAR"])
    add_list_validation(portfolio_ws, "J", ["Mantener", "Revisar precio", "Liquidar stock", "Reducir compra", "Discontinuar", "Mantener con monitoreo"])
    add_list_validation(portfolio_ws, "L", ["Alta", "Media", "Baja"])
    add_list_validation(portfolio_ws, "M", ["Alto", "Medio", "Bajo"])

    write_rows(
        workbook,
        "06_PRICING_DECISIONS",
        [
            ["sku_id", "sku_name", "category", "portfolio_decision", "current_price_m12", "avg_price_12m", "current_cost_m12", "gross_margin_pct", "elasticity_proxy", "pricing_decision", "price_m13", "price_m14", "price_m15", "expected_volume_effect", "rationale", "risk", "ai_comment", "team_comment"],
            ["Completar en Ejercicio 2", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
    )
    add_list_validation(workbook["06_PRICING_DECISIONS"], "J", ["Mantener precio", "Subir precio", "Bajar precio", "Precio promocional", "Liquidacion", "Revisar competitividad"])

    write_rows(
        workbook,
        "07_FORECAST_90_DIAS",
        [
            ["sku_id", "sku_name", "category", "portfolio_decision", "scenario", "baseline_units_m13", "baseline_units_m14", "baseline_units_m15", "baseline_revenue_m13", "baseline_revenue_m14", "baseline_revenue_m15", "projected_units_m13", "projected_units_m14", "projected_units_m15", "projected_revenue_m13", "projected_revenue_m14", "projected_revenue_m15", "projected_margin_m13", "projected_margin_m14", "projected_margin_m15", "assumption", "risk"],
            ["Completar en Ejercicio 3", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
    )
    add_list_validation(workbook["07_FORECAST_90_DIAS"], "E", ["Conservador", "Base", "Agresivo"])

    write_rows(
        workbook,
        "08_INVENTORY_DECISIONS",
        [
            ["sku_id", "sku_name", "category", "portfolio_decision", "forecast_units_90d", "current_stock", "current_inventory_value", "ddi_current", "ddi_target", "target_stock", "stock_gap", "inventory_action", "capital_released", "stockout_risk", "rationale", "ai_comment", "team_comment"],
            ["Completar en Ejercicio 4", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
    )
    add_list_validation(workbook["08_INVENTORY_DECISIONS"], "L", ["Mantener stock", "Reducir compra", "Liquidar excedente", "Reponer", "Bloquear reposicion", "Monitorear"])
    add_list_validation(workbook["08_INVENTORY_DECISIONS"], "N", ["Alto", "Medio", "Bajo"])

    write_rows(
        workbook,
        "09_SCOREBOARD_ALUMNO",
        [
            ["exercise", "metric", "initial_value", "proposed_value", "impact", "assumption", "source_or_formula"],
            ["Ejercicio 0", "Revenue historico", "", "", "", "", ""],
            ["Ejercicio 1", "SKUs Core", "", "", "", "", ""],
            ["Ejercicio 1", "SKUs Review", "", "", "", "", ""],
            ["Ejercicio 1", "SKUs Eliminar", "", "", "", "", ""],
            ["Ejercicio 1", "Capital liberado estimado", "", "", "", "", ""],
            ["Ejercicio 2", "Margen proyectado", "", "", "", "", ""],
            ["Ejercicio 3", "Revenue forecast M13-M15", "", "", "", "", ""],
            ["Ejercicio 4", "Capital de trabajo objetivo", "", "", "", "", ""],
        ],
    )
    write_rows(
        workbook,
        "10_PLAN_90_DIAS",
        [
            ["initiative_id", "initiative_name", "related_skus_or_category", "month_1_action", "month_2_action", "month_3_action", "expected_impact", "owner", "risk", "dependency"],
            ["P01", "", "", "", "", "", "", "", "", ""],
        ],
    )
    write_rows(
        workbook,
        "11_OUTPUT_FINAL",
        [
            ["campo", "respuesta_del_equipo"],
            ["tesis_final", ""],
            ["decisiones_adoptadas", ""],
            ["impacto_economico", ""],
            ["quick_wins", ""],
            ["roadmap_implementacion", ""],
            ["riesgos", ""],
            ["roi_esperado", ""],
        ],
    )
    write_rows(
        workbook,
        "12_CONTROL_STATUS",
        [
            ["exercise", "main_sheet", "status", "completed_by_team", "notes"],
            ["Ejercicio 0", "03_GUIA_EXPLORACION / 04_DIAGNOSTICO_INICIAL", "pending", "", ""],
            ["Ejercicio 1", "05_DECISIONES_PORTFOLIO", "pending", "", ""],
            ["Ejercicio 2", "06_PRICING_DECISIONS", "pending", "", ""],
            ["Ejercicio 3", "07_FORECAST_90_DIAS", "pending", "", ""],
            ["Ejercicio 4", "08_INVENTORY_DECISIONS", "pending", "", ""],
            ["Plan final", "11_OUTPUT_FINAL", "pending", "", ""],
        ],
    )

    ordered = {sheet.title: sheet for sheet in workbook.worksheets}
    workbook._sheets = [ordered[name] for name in EXPECTED_SHEETS]
    return workbook


def main() -> None:
    workbook = build_workbook()
    workbook.save(OUTPUT)
    APP_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(OUTPUT, APP_OUTPUT)
    print(f"Workbook generado: {OUTPUT}")
    print(f"Copia Next.js: {APP_OUTPUT}")


if __name__ == "__main__":
    main()
