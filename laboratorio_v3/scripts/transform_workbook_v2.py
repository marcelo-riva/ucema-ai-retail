from __future__ import annotations

import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]
INPUT = REPO_ROOT / "laboratorio_v3" / "public" / "templates" / "NEXUS_RETAIL_LAB01_WORKBOOK_COMPLETO.xlsx"
OUTPUT = REPO_ROOT / "laboratorio_v3" / "public" / "templates" / "NEXUS_RETAIL_LAB01_WORKBOOK_COMPLETO.xlsx"
APP_OUTPUT = REPO_ROOT / "laboratorio_v3" / "apps" / "web" / "public" / "templates" / "NEXUS_RETAIL_LAB01_WORKBOOK_COMPLETO.xlsx"
BACKUP = REPO_ROOT / "laboratorio_v3" / "public" / "templates" / "NEXUS_RETAIL_LAB01_WORKBOOK_COMPLETO_BACKUP.xlsx"

SHEET_RENAME = {
    "00_CASO_NEGOCIO": "01_CASO_NEGOCIO",
    "01_BASE_SKUS": "03_BASE_SKUS",
    "02_DICCIONARIO_DATOS": "02_DICCIONARIO_DATOS",
    "03_GUIA_EXPLORACION": "04_EXPLORACION",
    "04_DIAGNOSTICO_INICIAL": "05_DIAGNOSTICO_INICIAL",
    "05_DECISIONES_PORTFOLIO": "06_PORTFOLIO",
    "06_PRICING_DECISIONS": "09_PRICING",
    "07_FORECAST_90_DIAS": "07_FORECAST_90_DIAS",
    "08_INVENTORY_DECISIONS": "08_INVENTARIO",
    "09_SCOREBOARD_ALUMNO": "11_SCOREBOARD_ALUMNO",
    "10_PLAN_90_DIAS": "12_PLAN_90_DIAS",
    "11_OUTPUT_FINAL": "13_OUTPUT_FINAL",
    "12_CONTROL_STATUS": "14_CONTROL_STATUS",
    "13_LISTS": "15_LISTS",
}

FINAL_ORDER = [
    "00_INSTRUCCIONES",
    "01_CASO_NEGOCIO",
    "02_DICCIONARIO_DATOS",
    "03_BASE_SKUS",
    "04_EXPLORACION",
    "05_DIAGNOSTICO_INICIAL",
    "06_PORTFOLIO",
    "07_FORECAST_90_DIAS",
    "08_INVENTARIO",
    "09_PRICING",
    "11_SCOREBOARD_ALUMNO",
    "12_PLAN_90_DIAS",
    "13_OUTPUT_FINAL",
    "14_CONTROL_STATUS",
    "15_LISTS",
]

# Estilos
HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=16, color="0F172A")
SUBTITLE_FONT = Font(bold=True, size=12, color="0F172A")
EDITABLE_FILL = PatternFill("solid", fgColor="FEF9C3")
FORMULA_FILL = PatternFill("solid", fgColor="F1F5F9")
REF_FILL = PatternFill("solid", fgColor="E0F2FE")
GUIDE_FILL = PatternFill("solid", fgColor="ECFDF5")
NOTE_FILL = PatternFill("solid", fgColor="FEF3C7")
WHITE_FONT = Font(color="FFFFFF", bold=True)
BOLD_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

# Hojas que reciben el bloque de guia arriba (5 filas)
GUIDE_SHEETS = ["04_EXPLORACION", "05_DIAGNOSTICO_INICIAL", "06_PORTFOLIO",
                "07_FORECAST_90_DIAS", "08_INVENTARIO", "09_PRICING",
                "12_PLAN_90_DIAS", "13_OUTPUT_FINAL"]
SHEET_ROW_SHIFTS = {sheet: 5 for sheet in GUIDE_SHEETS}

# Configuracion de ejercicios para scoreboard y control
EXERCISE_CONFIG = [
    # (exercise, sheet, key_col, key_start_row, min_required, next_action)
    ("Ejercicio 0", "04_EXPLORACION", "D", 7, 5, "Completar hallazgos e implicancias clave."),
    ("Ejercicio 0", "05_DIAGNOSTICO_INICIAL", "B", 8, 5, "Completar diagnóstico inicial y preguntas de negocio."),
    ("Ejercicio 1", "06_PORTFOLIO", "O", 7, 10, "Clasificar al menos 10 SKUs en CORE / REVIEW / ELIMINAR."),
    ("Ejercicio 2", "07_FORECAST_90_DIAS", "E", 7, 10, "Completar proyección de al menos 10 SKUs."),
    ("Ejercicio 3", "08_INVENTARIO", "I", 7, 10, "Definir inventario de al menos 10 SKUs."),
    ("Ejercicio 4", "09_PRICING", "N", 7, 10, "Definir pricing de al menos 10 SKUs."),
    ("Plan", "12_PLAN_90_DIAS", "C", 7, 3, "Completar al menos 3 iniciativas con owner y estado."),
    ("Cierre", "13_OUTPUT_FINAL", "B", 8, 5, "Completar tesis y resumen ejecutivo."),
]


def apply_header_style(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def set_column_widths(ws, min_width: int = 14, max_width: int = 40, sample_rows: int = 100) -> None:
    max_row = min(sample_rows, ws.max_row)
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_length = 0
        for row in range(1, max_row + 1):
            cell = ws.cell(row, col_idx)
            try:
                length = len(str(cell.value or ""))
                if length > max_length:
                    max_length = length
            except Exception:
                pass
        width = min(max(max_length + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = width


def fill_range(ws, col_letter: str, start_row: int, end_row: int, fill: PatternFill) -> None:
    for row in range(start_row, end_row + 1):
        ws[f"{col_letter}{row}"].fill = fill


def shift_formula_references(formula: str, source_sheet: str, sheet_shifts: dict[str, int]) -> str:
    """Suma el desplazamiento de filas a las referencias de celda de una formula.

    Solo desplaza referencias a filas >= 2 (zona de datos original), no encabezados.
    """
    # Patron que captura referencias como A1, 'Sheet'!A1, Sheet!A1, $A$1, A1:B2, etc.
    pattern = re.compile(
        r"((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(\$?[A-Z]{1,3})(\$?\d+)"
        r"(?::(\$?[A-Z]{1,3})(\$?\d+))?"
    )

    def replace(match):
        sheet_ref = match.group(1)
        col1 = match.group(2)
        row1_str = match.group(3)
        col2 = match.group(4)
        row2_str = match.group(5)

        target_sheet = sheet_ref.rstrip("!").replace("'", "") if sheet_ref else source_sheet
        if target_sheet in sheet_shifts:
            shift = sheet_shifts[target_sheet]
            row1 = int(row1_str.replace("$", ""))
            if row1 >= 2:
                new_row1 = row1 + shift
                row1_str = f"${new_row1}" if "$" in row1_str else str(new_row1)
            if row2_str:
                row2 = int(row2_str.replace("$", ""))
                if row2 >= 2:
                    new_row2 = row2 + shift
                    row2_str = f"${new_row2}" if "$" in row2_str else str(new_row2)

        range_part = f":{col2}{row2_str}" if row2_str else ""
        return f"{sheet_ref or ''}{col1}{row1_str}{range_part}"

    return pattern.sub(replace, formula)


def update_all_formula_references(wb) -> None:
    """Actualiza nombres de hojas y desplazamientos de filas en todas las formulas."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    new_formula = cell.value
                    # Renombrar hojas
                    for old_name, new_name in SHEET_RENAME.items():
                        new_formula = new_formula.replace(f"'{old_name}'!", f"'{new_name}'!")
                        new_formula = new_formula.replace(f"{old_name}!", f"{new_name}!")
                    # Desplazar filas de hojas que insertaron guia
                    new_formula = shift_formula_references(new_formula, ws.title, SHEET_ROW_SHIFTS)
                    if new_formula != cell.value:
                        cell.value = new_formula


def insert_guide_block(ws, title: str, objective: str, excel: str, platform: str, completion: str) -> None:
    ws.insert_rows(1, 5)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A1"].fill = GUIDE_FILL
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ws.max_column)
    ws["A2"] = f"Objetivo: {objective}"
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A2"].fill = GUIDE_FILL
    ws.row_dimensions[2].height = 25

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ws.max_column)
    ws["A3"] = f"Excel: {excel}"
    ws["A3"].font = BOLD_FONT
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A3"].fill = GUIDE_FILL
    ws.row_dimensions[3].height = 25

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=ws.max_column)
    ws["A4"] = f"Plataforma: {platform}"
    ws["A4"].font = BOLD_FONT
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A4"].fill = GUIDE_FILL
    ws.row_dimensions[4].height = 25

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=ws.max_column)
    ws["A5"] = f"Criterio de completitud: {completion}"
    ws["A5"].font = BOLD_FONT
    ws["A5"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A5"].fill = GUIDE_FILL
    ws.row_dimensions[5].height = 25


def create_instructions_sheet(wb) -> None:
    ws = wb.create_sheet("00_INSTRUCCIONES", 0)

    ws["A1"] = "NEXUS RETAIL — LABORATORIO 1: AI REVENUE & INVENTORY COPILOT"
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, size=18, color="0F172A")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws["A3"] = "Este workbook es el espacio de trabajo numérico del grupo. Usalo para analizar SKUs, forecast, inventario, pricing y decisiones operativas. La plataforma web guía el laboratorio y captura los hallazgos cualitativos, decisiones, riesgos y plan final. No copies en la plataforma toda la información línea por línea: llevá solo la síntesis que defenderías frente a un equipo comercial. El Scoreboard se actualiza automáticamente a medida que completás las hojas de trabajo."
    ws.merge_cells("A3:D3")
    ws["A3"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws["A3"].font = Font(size=11)
    ws.row_dimensions[3].height = 100

    ws["A5"] = "¿Qué se completa en Excel?"
    ws["A5"].font = SUBTITLE_FONT
    ws["A6"] = "• Análisis numérico, simulación, datos por SKU, forecast, inventario, pricing y decisiones operativas."
    ws["A7"] = "• Cálculos, escenarios y jugadas por producto."
    ws["A8"] = "• Datos de referencia y backup operativo."

    ws["A10"] = "¿Qué se registra en la plataforma?"
    ws["A10"].font = SUBTITLE_FONT
    ws["A11"] = "• Hallazgos principales, criterio de decisión, riesgos y trade-offs."
    ws["A12"] = "• Síntesis ejecutiva y plan final narrativo."
    ws["A13"] = "• Supuestos, justificaciones y dependencias."

    ws["A15"] = "Leyenda de celdas"
    ws["A15"].font = SUBTITLE_FONT
    ws["A16"] = "Celdas editables por el alumno"
    ws["A16"].fill = EDITABLE_FILL
    ws["B16"] = "Fondo amarillo suave. Campos que el equipo debe completar."
    ws["A17"] = "Fórmulas / cálculos automáticos"
    ws["A17"].fill = FORMULA_FILL
    ws["B17"] = "Fondo gris claro. No editar."
    ws["A18"] = "Referencias / no editar"
    ws["A18"].fill = REF_FILL
    ws["B18"] = "Fondo azul suave. Datos traidos desde otras hojas."
    ws["A19"] = "Encabezados"
    ws["A19"].fill = HEADER_FILL
    ws["A19"].font = WHITE_FONT
    ws["B19"] = "Fondo oscuro con letra blanca."

    ws["A21"] = "Orden de trabajo sugerido"
    ws["A21"].font = SUBTITLE_FONT
    order_text = [
        "1. Leer 01_CASO_NEGOCIO y 02_DICCIONARIO_DATOS.",
        "2. Explorar 03_BASE_SKUS y completar 04_EXPLORACION + 05_DIAGNOSTICO_INICIAL.",
        "3. Decidir portfolio en 06_PORTFOLIO.",
        "4. Proyectar 90 días en 07_FORECAST_90_DIAS.",
        "5. Definir inventario en 08_INVENTARIO.",
        "6. Definir pricing en 09_PRICING.",
        "7. Revisar 11_SCOREBOARD_ALUMNO (automático).",
        "8. Cerrar plan en 12_PLAN_90_DIAS y 13_OUTPUT_FINAL.",
    ]
    for i, line in enumerate(order_text, start=22):
        ws[f"A{i}"] = line

    ws["A31"] = "Nota importante"
    ws["A31"].font = SUBTITLE_FONT
    ws["A32"] = "El Scoreboard es automático. No hace falta completarlo manualmente. Las hojas 14_CONTROL_STATUS y 15_LISTS son técnicas y quedan al final."
    ws.merge_cells("A32:D32")
    ws["A32"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws["A32"].fill = NOTE_FILL
    ws.row_dimensions[32].height = 45

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20


def build_scoreboard(wb) -> None:
    # Reemplazar la hoja de scoreboard por una nueva
    idx = wb.sheetnames.index("11_SCOREBOARD_ALUMNO")
    del wb["11_SCOREBOARD_ALUMNO"]
    ws = wb.create_sheet("11_SCOREBOARD_ALUMNO", idx)

    ws["A1"] = "SCOREBOARD ALUMNO — Se actualiza automáticamente"
    ws.merge_cells("A1:G1")
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws["A1"].fill = NOTE_FILL
    ws.row_dimensions[1].height = 30

    ws["A2"] = "No hace falta completar esta hoja manualmente. Lee el estado de avance de cada ejercicio."
    ws.merge_cells("A2:G2")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A2"].fill = NOTE_FILL
    ws.row_dimensions[2].height = 25

    headers = ["Ejercicio", "Hoja", "Estado", "Completos", "Mínimo requerido", "% completitud", "Próxima acción sugerida"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(4, c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for i, (exercise, sheet, key_col, key_start_row, min_required, next_action) in enumerate(EXERCISE_CONFIG, start=5):
        row = i
        ws.cell(row, 1, value=exercise)
        ws.cell(row, 2, value=sheet)
        ws.cell(row, 3, value=f'=IF(COUNTA(\'{sheet}\'!{key_col}{key_start_row}:{key_col}8226)=0,"Pendiente",IF(COUNTA(\'{sheet}\'!{key_col}{key_start_row}:{key_col}8226)<{min_required},"En progreso","Completo"))')
        ws.cell(row, 4, value=f"=COUNTA('{sheet}'!{key_col}{key_start_row}:{key_col}8226)")
        ws.cell(row, 5, value=min_required)
        ws.cell(row, 6, value=f"=IF(E{row}=0,0,D{row}/E{row})")
        ws.cell(row, 6).number_format = "0%"
        ws.cell(row, 7, value=next_action)

        for c in range(1, 8):
            ws.cell(row, c).border = THIN_BORDER
            ws.cell(row, c).alignment = Alignment(vertical="center", wrap_text=True)

    total_row = 5 + len(EXERCISE_CONFIG)
    ws.cell(total_row, 1, value="Avance total")
    ws.cell(total_row, 1).font = BOLD_FONT
    ws.cell(total_row, 6, value=f"=AVERAGE(F5:F{total_row-1})")
    ws.cell(total_row, 6).number_format = "0%"
    ws.cell(total_row, 6).font = BOLD_FONT
    for c in range(1, 8):
        ws.cell(total_row, c).border = THIN_BORDER

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 45


def style_large_sheet(ws, editable_cols: list[int] = None, formula_cols: list[int] = None, reference_cols: list[int] = None) -> None:
    editable_cols = editable_cols or []
    formula_cols = formula_cols or []
    reference_cols = reference_cols or []
    max_row = ws.max_row
    for col in editable_cols:
        fill_range(ws, get_column_letter(col), 7, max_row, EDITABLE_FILL)
    for col in formula_cols:
        fill_range(ws, get_column_letter(col), 7, max_row, FORMULA_FILL)
    for col in reference_cols:
        fill_range(ws, get_column_letter(col), 7, max_row, REF_FILL)


def style_small_sheet(ws, editable_cols: list[int] = None, formula_cols: list[int] = None, reference_cols: list[int] = None, max_rows: int = 200) -> None:
    editable_cols = editable_cols or []
    formula_cols = formula_cols or []
    reference_cols = reference_cols or []
    for row in range(2, min(ws.max_row + 1, max_rows + 1)):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row, c)
            if c in editable_cols:
                cell.fill = EDITABLE_FILL
            elif c in formula_cols:
                cell.fill = FORMULA_FILL
            elif c in reference_cols:
                cell.fill = REF_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def add_filters(ws, header_row: int = 6) -> None:
    if ws.max_row > header_row and ws.max_column > 1:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"


def setup_control_status(ws) -> None:
    """Configura 14_CONTROL_STATUS con columnas de control y umbrales."""
    # Limpiar y poner headers
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
            cell.fill = PatternFill()
            cell.font = Font()

    headers = ["exercise_id", "exercise_name", "main_sheet", "status", "completed_by_team",
               "platform_validation", "key_column", "key_start_row", "min_required", "comments"]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, value=h)

    for i, (exercise, sheet, key_col, key_start_row, min_required, _) in enumerate(EXERCISE_CONFIG, start=2):
        ws.cell(i, 1, value=f"EX0{i-2:02d}" if "Ejercicio" in exercise else ("PLAN" if "Plan" in exercise else "CLOS"))
        ws.cell(i, 2, value=exercise)
        ws.cell(i, 3, value=sheet)
        ws.cell(i, 4, value="PENDING")
        ws.cell(i, 5, value="NO")
        ws.cell(i, 6, value="NOT_VALIDATED")
        ws.cell(i, 7, value=key_col)
        ws.cell(i, 8, value=key_start_row)
        ws.cell(i, 9, value=min_required)
        ws.cell(i, 10, value=f"Completar al menos {min_required} filas/campos en {sheet}.")

    apply_header_style(ws)
    set_column_widths(ws, sample_rows=10)
    style_small_sheet(ws)


def main() -> None:
    shutil.copyfile(INPUT, BACKUP)
    print(f"Backup guardado en: {BACKUP}")

    wb = load_workbook(INPUT, data_only=False)

    # Renombrar hojas
    for old_name, new_name in SHEET_RENAME.items():
        if old_name in wb.sheetnames:
            wb[old_name].title = new_name
            print(f"Renombrada: {old_name} -> {new_name}")

    # Insertar bloques de guia en hojas de ejercicio (antes de actualizar formulas)
    for sheet_name in GUIDE_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        if sheet_name == "04_EXPLORACION":
            insert_guide_block(ws, "Ejercicio 0 — Exploración Inicial",
                               "Entender la base, las variables y las señales antes de tomar decisiones.",
                               "Análisis de la base de SKUs, respuestas a las preguntas guía y hallazgos numéricos.",
                               "Diagnóstico cualitativo, interpretación de hallazgos, riesgos y síntesis ejecutiva.",
                               "Completar las columnas de hallazgo e implicancia para cada pregunta.")
        elif sheet_name == "05_DIAGNOSTICO_INICIAL":
            insert_guide_block(ws, "Ejercicio 0 — Diagnóstico Inicial",
                               "Sintetizar el diagnóstico del negocio y las preguntas clave que guiarán las decisiones.",
                               "Resumen numérico del negocio y señales encontradas en la exploración.",
                               "Diagnóstico narrativo, hallazgos principales, preguntas de negocio y supuestos.",
                               "Completar todos los campos de respuesta del equipo.")
        elif sheet_name == "06_PORTFOLIO":
            insert_guide_block(ws, "Ejercicio 1 — Portfolio Optimization",
                               "Clasificar cada SKU en CORE, REVIEW o ELIMINAR y definir la acción a 90 días.",
                               "Análisis por SKU, clasificación de portfolio, acción, prioridad y riesgo comercial.",
                               "Criterio de clasificación, hallazgos por categoría, riesgos y síntesis de decisión.",
                               "Completar decisión de portfolio, acción, prioridad y riesgo para los SKUs relevantes.")
        elif sheet_name == "07_FORECAST_90_DIAS":
            insert_guide_block(ws, "Ejercicio 2 — Forecast Engine",
                               "Proyectar unidades, precios, costos y stock para M13-M15, y estimar revenue y margen.",
                               "Proyección numérica por SKU, escenario, revenue y margen proyectado.",
                               "Supuestos del forecast, riesgos y síntesis de escenarios.",
                               "Completar escenario, unidades proyectadas, precios, costos y supuestos.")
        elif sheet_name == "08_INVENTARIO":
            insert_guide_block(ws, "Ejercicio 3 — Inventory & Working Capital",
                               "Definir stock objetivo, DDI objetivo y acción de inventario para liberar capital o evitar quiebres.",
                               "Cálculo de DDI objetivo, stock gap, capital liberado y riesgo de quiebre.",
                               "Riesgos de inventario, trade-offs capital-servicio y síntesis de acciones.",
                               "Completar DDI objetivo, stock objetivo, acción de inventario, capital liberado y riesgo.")
        elif sheet_name == "09_PRICING":
            insert_guide_block(ws, "Ejercicio 4 — Pricing Optimization",
                               "Definir la decisión de precio por SKU y los precios M13-M15, estimando el efecto en volumen.",
                               "Análisis de precio actual, costo, margen, competencia y decisión de precio.",
                               "Justificación de pricing, riesgos y síntesis de jugadas de precio.",
                               "Completar decisión de pricing, precios M13-M15, efecto esperado y riesgo.")
        elif sheet_name == "12_PLAN_90_DIAS":
            insert_guide_block(ws, "Plan de captura de valor — 90 días",
                               "Traducir las decisiones en iniciativas concretas con acciones a 30, 60 y 90 días.",
                               "Resumen de iniciativas, impactos esperados y responsables.",
                               "Narrativa del plan, dependencias y riesgos de ejecución.",
                               "Completar al menos 3 iniciativas con impacto, owner y estado.")
        elif sheet_name == "13_OUTPUT_FINAL":
            insert_guide_block(ws, "Output final — Tesis y resumen ejecutivo",
                               "Cerrar el laboratorio con una tesis clara, decisiones adoptadas e impacto esperado.",
                               "Resumen de números clave y decisiones tomadas.",
                               "Tesis final, riesgos, quick wins y roadmap de implementación.",
                               "Completar los campos de resumen ejecutivo.")
        print(f"Guía insertada en {sheet_name}")

    # Crear hoja de instrucciones
    create_instructions_sheet(wb)

    # Actualizar todas las referencias de formulas (nombres de hoja y desplazamientos de fila)
    print("Actualizando referencias de formulas...")
    update_all_formula_references(wb)
    print("Referencias actualizadas.")

    # Reordenar hojas
    ordered = {ws.title: ws for ws in wb.worksheets}
    wb._sheets = [ordered[name] for name in FINAL_ORDER if name in ordered]

    # Aplicar estilos y filtros a hojas de ejercicio
    # 04_EXPLORACION
    ws = wb["04_EXPLORACION"]
    apply_header_style(ws, row=6)
    add_filters(ws)
    set_column_widths(ws, sample_rows=20)
    style_small_sheet(ws, editable_cols=[4, 5])
    ws.freeze_panes = "A7"

    # 05_DIAGNOSTICO_INICIAL
    ws = wb["05_DIAGNOSTICO_INICIAL"]
    set_column_widths(ws, sample_rows=30)
    style_small_sheet(ws, editable_cols=[2])

    # 06_PORTFOLIO
    ws = wb["06_PORTFOLIO"]
    apply_header_style(ws, row=6)
    add_filters(ws)
    set_column_widths(ws, sample_rows=10)
    style_large_sheet(ws, editable_cols=[15, 16, 17, 18, 19, 20, 21], formula_cols=[6, 7, 8, 9, 10, 11, 12, 13, 14])
    ws.freeze_panes = "G7"

    # 07_FORECAST_90_DIAS
    ws = wb["07_FORECAST_90_DIAS"]
    apply_header_style(ws, row=6)
    add_filters(ws)
    set_column_widths(ws, sample_rows=10)
    style_large_sheet(ws, editable_cols=[6, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 36, 37, 38, 39],
                      formula_cols=[30, 31, 32, 33, 34, 35], reference_cols=[1, 2, 3, 4, 5])
    ws.freeze_panes = "G7"

    # 08_INVENTARIO
    ws = wb["08_INVENTARIO"]
    apply_header_style(ws, row=6)
    add_filters(ws)
    set_column_widths(ws, sample_rows=10)
    style_large_sheet(ws, editable_cols=[9, 10, 12, 13, 14, 15, 16, 17], formula_cols=[5, 11], reference_cols=[1, 2, 3, 4, 6, 7, 8])
    ws.freeze_panes = "G7"

    # 09_PRICING
    ws = wb["09_PRICING"]
    apply_header_style(ws, row=6)
    add_filters(ws)
    set_column_widths(ws, sample_rows=10)
    style_large_sheet(ws, editable_cols=[14, 15, 16, 17, 18, 19, 20, 21, 22], formula_cols=[10, 11, 12, 13], reference_cols=[1, 2, 3, 4, 5, 6, 7, 8, 9])
    ws.freeze_panes = "G7"

    # 12_PLAN_90_DIAS
    ws = wb["12_PLAN_90_DIAS"]
    apply_header_style(ws, row=6)
    add_filters(ws)
    set_column_widths(ws, sample_rows=30)
    style_small_sheet(ws, editable_cols=[2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14])
    ws.freeze_panes = "A7"

    # 13_OUTPUT_FINAL
    ws = wb["13_OUTPUT_FINAL"]
    set_column_widths(ws, sample_rows=20)
    style_small_sheet(ws, editable_cols=[2])

    # 03_BASE_SKUS: no mover filas, A1 = sku_id, agregar comentario
    ws = wb["03_BASE_SKUS"]
    ws["A1"].comment = Comment("NO EDITAR: esta hoja conserva la base histórica M01-M12 del Laboratorio 1.", "NEXUS Retail")
    apply_header_style(ws, row=1)
    ws.freeze_panes = "A2"
    set_column_widths(ws, sample_rows=10)
    fill_range(ws, "A", 2, ws.max_row, REF_FILL)

    # 02_DICCIONARIO_DATOS
    ws = wb["02_DICCIONARIO_DATOS"]
    apply_header_style(ws)
    set_column_widths(ws, sample_rows=12)
    style_small_sheet(ws)

    # 01_CASO_NEGOCIO
    ws = wb["01_CASO_NEGOCIO"]
    ws["A1"] = "NEXUS RETAIL — CASO DE NEGOCIO"
    ws.merge_cells("A1:D1")
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 30
    set_column_widths(ws, sample_rows=23)

    # 14_CONTROL_STATUS
    ws = wb["14_CONTROL_STATUS"]
    setup_control_status(ws)

    # 15_LISTS
    ws = wb["15_LISTS"]
    apply_header_style(ws)
    set_column_widths(ws, sample_rows=7)
    style_small_sheet(ws)

    # Scoreboard
    print("Construyendo scoreboard...")
    build_scoreboard(wb)
    print("Scoreboard listo.")

    # Guardar
    wb.save(OUTPUT)
    APP_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(OUTPUT, APP_OUTPUT)
    print(f"Workbook transformado guardado en: {OUTPUT}")
    print(f"Copia Next.js: {APP_OUTPUT}")


if __name__ == "__main__":
    main()
